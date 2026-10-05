"""Bulk-load an existing mortgage book from a spreadsheet (CSV or Excel)."""
from __future__ import annotations

import io
import re
from datetime import date, datetime
from decimal import Decimal

from ..models import PROPERTY_TYPES, STATUSES
from . import calc
from . import statements as st

# (header, form field, help)
COLUMNS = [
    ("Reference", "reference", "Optional — e.g. M-001. Generated if blank."),
    ("Borrower name", "borrower_name", "Required"),
    ("Borrower email", "borrower_email", ""),
    ("Borrower phone", "borrower_phone", ""),
    ("Property address", "property_address", "Required"),
    ("City", "property_city", ""),
    ("Province", "property_province", "e.g. ON"),
    ("Property type", "property_type", "Detached, Semi-detached, Townhouse, Condo, Multi-unit, Commercial, Mixed use, Land, Other"),
    ("Property value", "property_value", "Appraised value"),
    ("Appraisal date", "appraisal_date", "YYYY-MM-DD"),
    ("Position", "position", "1, 2 or 3"),
    ("Prior charges", "prior_charges", "Balance of mortgages ranking ahead of yours"),
    ("Principal", "principal_amount", "Required — amount advanced"),
    ("Rate %", "interest_rate", "Required — e.g. 10.99"),
    ("Rate type", "rate_type", "Fixed or Variable"),
    ("Prime spread %", "prime_spread", "Variable only — e.g. 5.0 for prime + 5%"),
    ("Rate floor %", "rate_floor", "Variable only"),
    ("Compounding", "compounding", "Monthly, Semi-annual or Annual"),
    ("Payment type", "payment_type", "Interest only or Amortizing"),
    ("Payment frequency", "payment_frequency", "Monthly, Semi-monthly, Bi-weekly, Weekly, Quarterly, Annually"),
    ("Payment amount", "payment_amount", "Leave blank to calculate"),
    ("Amortization months", "amortization_months", "Amortizing only, e.g. 300"),
    ("Funded date", "funded_date", "Required — YYYY-MM-DD"),
    ("First payment date", "first_payment_date", "Blank = one month after funding"),
    ("Term months", "term_months", "Term or maturity date required"),
    ("Maturity date", "maturity_date", "YYYY-MM-DD"),
    ("Your share %", "ownership_pct", "100 unless syndicated"),
    ("Owners", "owners", "Who owns or funded it, e.g. 9929916 Canada Inc"),
    ("Lender fee", "lender_fee", ""),
    ("Broker", "broker_name", ""),
    ("Broker fee", "broker_fee", ""),
    ("Lawyer", "lawyer_name", ""),
    ("Renewal fee", "renewal_fee", ""),
    ("NSF fee", "nsf_fee", ""),
    ("Insurance expiry", "insurance_expiry", "YYYY-MM-DD"),
    ("Status", "status", "Active, In arrears, Default, Matured, Paid out"),
    ("Bank keywords", "match_keywords", "Text on your bank statement when they pay; separate several with ;"),
    ("Notes", "notes", ""),
    ("Current balance", "_balance", "Optional — if principal has been repaid, the balance today"),
    ("Balance as of", "_balance_date", "Date of that balance (default today)"),
]
DATE_FIELDS = {"appraisal_date", "funded_date", "first_payment_date", "maturity_date", "insurance_expiry", "_balance_date"}

def _key(text) -> str:
    return re.sub(r"[^a-z0-9]", "", str(text or "").lower())


def _choice_map(pairs, extra=None):
    out = {}
    for value, label in pairs:
        out[_key(value)] = value
        out[_key(label)] = value
    out.update({_key(k): v for k, v in (extra or {}).items()})
    return out


CHOICE_MAPS = {
    "property_type": _choice_map(PROPERTY_TYPES, {"semi": "semi", "multi-unit": "multi_unit", "multiunit": "multi_unit",
                                                  "land": "land", "house": "detached", "town house": "townhouse"}),
    "status": _choice_map(STATUSES, {"default": "default", "paid": "paid_out", "arrears": "in_arrears"}),
    "rate_type": {"fixed": "fixed", "variable": "variable", "floating": "variable", "prime": "variable"},
    "compounding": {"monthly": "monthly", "semiannual": "semi_annual", "semiannually": "semi_annual",
                    "annual": "annual", "annually": "annual", "yearly": "annual"},
    "payment_type": {"interestonly": "interest_only", "io": "interest_only", "amortizing": "amortizing",
                     "amortized": "amortizing", "blended": "amortizing", "pi": "amortizing"},
    "payment_frequency": {_key(k): k for k in calc.FREQUENCIES} | {"semimonthly": "semi_monthly", "biweekly": "biweekly",
                                                                  "fortnightly": "biweekly", "yearly": "annually",
                                                                  "annual": "annually"},
}


def template_workbook() -> bytes:
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Mortgages"
    ws.append([h for h, _, _ in COLUMNS])
    for i, (h, _field, help_text) in enumerate(COLUMNS, start=1):
        c = ws.cell(row=1, column=i)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="14213D" if "Required" not in help_text else "9A031E")
        if help_text:
            c.comment = Comment(help_text, "Mortgage Manager")
        ws.column_dimensions[get_column_letter(i)].width = max(14, len(h) + 4)
    ws.append(["", "Jane Smith", "jane@example.com", "416-555-0100", "12 Maple Ave", "Toronto", "ON", "Detached",
               1150000, "2026-01-10", 2, 610000, 250000, 10.99, "Fixed", "", "", "Monthly", "Interest only", "Monthly",
               "", "", "2026-02-01", "2026-03-01", 12, "", 100, "Example Holdings Inc", 5000, "Example Brokerage", 2500, "", 1000, 300,
               "2027-01-31", "Active", "E-TRANSFER JANE SMITH", "Example row — delete before importing", "", ""])
    ws.freeze_panes = "A2"
    help_ws = wb.create_sheet("Instructions")
    help_ws.append(["Column", "What to enter"])
    for h, _, help_text in COLUMNS:
        help_ws.append([h, help_text])
    help_ws.column_dimensions["A"].width = 22
    help_ws.column_dimensions["B"].width = 90
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _to_form_value(field, raw):
    if raw is None:
        return ""
    if isinstance(raw, (datetime, date)):
        return (raw.date() if isinstance(raw, datetime) else raw).isoformat()
    text = str(raw).strip()
    if text == "":
        return ""
    if field in DATE_FIELDS:
        d = st.parse_date(text)
        return d.isoformat() if d else text
    if field in CHOICE_MAPS:
        return CHOICE_MAPS[field].get(_key(text), text)
    if field == "match_keywords":
        return "\n".join(k.strip() for k in re.split(r"[;\n]", text) if k.strip())
    if field == "position":
        m = re.match(r"(\d)", text)
        return m.group(1) if m else text
    if isinstance(raw, float) and raw.is_integer() and field in ("position", "term_months", "amortization_months"):
        return str(int(raw))
    return text


def parse_rows(filename: str, content: bytes):
    """[(row_number, form dict)] from the uploaded sheet; raises StatementError for unreadable files."""
    rows = st.read_rows(filename, content)
    header = [_key(h) for h in rows[0]]
    lookup = {_key(h): field for h, field, _ in COLUMNS}
    lookup.update({_key(field): field for _, field, _ in COLUMNS})
    cols = {i: lookup[h] for i, h in enumerate(header) if h in lookup}
    if "borrower_name" not in cols.values() or "principal_amount" not in cols.values():
        raise st.StatementError("The first row must be the template's column headers (Borrower name, Principal, …).")
    out = []
    for n, row in enumerate(rows[1:], start=2):
        form = {field: _to_form_value(field, row[i] if i < len(row) else None) for i, field in cols.items()}
        if not any(v for v in form.values()):
            continue
        if "example row" in (form.get("notes") or "").lower():
            continue
        out.append((n, form))
    return out


def balance_adjustment(form):
    """(as_of date, balance Decimal) when a current balance was given, else None."""
    raw = (form.get("_balance") or "").replace(",", "").replace("$", "").strip()
    if not raw:
        return None
    try:
        bal = Decimal(raw)
    except ArithmeticError:
        raise ValueError(f"Current balance '{raw}' is not a number") from None
    as_of = st.parse_date(form.get("_balance_date")) if form.get("_balance_date") else None
    return as_of, calc.money(bal)
