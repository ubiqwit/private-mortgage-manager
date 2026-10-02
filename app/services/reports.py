"""Month-end income report for the accountant (on screen, Excel and CSV)."""
from __future__ import annotations

import calendar
import csv
import io
from datetime import date, timedelta
from decimal import Decimal

from dateutil.relativedelta import relativedelta

from ..models import BankTransaction, Mortgage, MortgageTransaction
from ..timeutil import now as local_now
from ..timeutil import today as local_today
from . import calc
from .safety import csv_safe, set_text_cell

ZERO = Decimal("0.00")


def month_range(year: int, month: int):
    start = date(year, month, 1)
    return start, date(year, month, calendar.monthrange(year, month)[1])


def default_month(today: date | None = None):
    """The month just ended — what you'd hand the accountant."""
    today = today or local_today()
    prev = today.replace(day=1) - timedelta(days=1)
    return prev.year, prev.month


def accrued_interest(m: Mortgage, start: date, end: date, prime=None) -> Decimal:
    """Interest earned over [start, end] on the daily balance.

    Each day earns 1/days-in-month of the loan's monthly periodic rate, so a full month
    at a constant balance accrues exactly one month's contractual interest.
    """
    if end < m.funded_date:
        return ZERO
    total = Decimal(0)
    day = max(start, m.funded_date)
    while day <= end:
        days_in_month = calendar.monthrange(day.year, day.month)[1]
        # Interest runs on the balance at the start of the day; the funded day earns nothing.
        if day > m.funded_date:
            bal = m.balance(as_of=day - timedelta(days=1))
            if bal <= 0:
                break
            monthly = calc.periodic_rate(m.effective_rate(prime, on=day), m.terms_on(day).compounding, "monthly")
            total += bal * monthly / days_in_month
        day += timedelta(days=1)
    return calc.money(total)


def month_end_report(year: int, month: int, prime=None) -> dict:
    start, end = month_range(year, month)
    before = start - timedelta(days=1)

    mortgages = Mortgage.query.order_by(Mortgage.reference).all()
    rows = []
    for m in mortgages:
        txns = [t for t in m.transactions if start <= t.date <= end]
        opening = m.balance(as_of=before) if m.funded_date <= before else ZERO
        closing = m.balance(as_of=end) if m.funded_date <= end else ZERO
        active_in_month = (m.funded_date <= end) and (opening > 0 or closing > 0 or txns)
        if not active_in_month:
            continue
        inc = {"interest": ZERO, "principal": ZERO, "fees": ZERO, "advances": ZERO, "received": ZERO}
        for t in txns:
            if t.type == "advance":
                inc["advances"] += Decimal(str(t.amount))
                continue
            if t.type == "funding":
                continue  # the original principal is counted from the loan itself below
            inc["interest"] += Decimal(str(t.interest or 0))
            inc["principal"] += Decimal(str(t.principal or 0))
            inc["fees"] += Decimal(str(t.fees or 0))
            inc["received"] += t.signed_amount
        # A loan funded this month: its opening balance is zero and the advance is new money out.
        if m.funded_date >= start and m.funded_date <= end:
            inc["advances"] += Decimal(str(m.principal_amount))
        due = m.due_dates(start=start, end=end)
        expected = m.scheduled_total(start=start, end=end, prime=prime) if closing > 0 or opening > 0 else ZERO
        regular = m.regular_received(end) - m.regular_received(before)
        share = Decimal(str(m.ownership_pct or 100)) / 100
        accrued = accrued_interest(m, start, end, prime)
        rows.append(dict(
            m=m, opening=opening, closing=closing, rate=m.effective_rate(prime, on=end),
            advances=calc.money(inc["advances"]), principal=calc.money(inc["principal"]),
            interest=calc.money(inc["interest"]), fees=calc.money(inc["fees"]), received=calc.money(inc["received"]),
            income=calc.money(inc["interest"] + inc["fees"]),
            share_pct=Decimal(str(m.ownership_pct or 100)),
            your_income=calc.money((inc["interest"] + inc["fees"]) * share),
            accrued=accrued, expected=calc.money(expected), regular=calc.money(regular),
            variance=calc.money(regular - expected), payments_due=len(due),
            arrears=m.arrears(end, prime), status=m.display_status_label if end >= local_today() else _status_at(m, end),
        ))

    def total(key):
        return calc.money(sum((r[key] for r in rows), ZERO))

    totals = {k: total(k) for k in ("opening", "closing", "advances", "principal", "interest", "fees", "received",
                                    "income", "your_income", "accrued", "expected", "regular", "variance", "arrears")}

    transactions = (MortgageTransaction.query.filter(MortgageTransaction.date.between(start, end))
                    .order_by(MortgageTransaction.date, MortgageTransaction.id).all())

    deposits = BankTransaction.query.filter(BankTransaction.date.between(start, end), BankTransaction.amount > 0).all()
    bank = {
        "count": len(deposits),
        "total": calc.money(sum((Decimal(str(b.amount)) for b in deposits), ZERO)),
        "matched": calc.money(sum((b.allocated for b in deposits), ZERO)),
        "ignored": calc.money(sum((Decimal(str(b.amount)) for b in deposits if b.status == "ignored"), ZERO)),
        "unmatched_lines": sorted((b for b in deposits if b.status == "unmatched"), key=lambda b: b.date),
    }
    bank["unmatched"] = calc.money(sum((Decimal(str(b.amount)) - b.allocated for b in bank["unmatched_lines"]), ZERO))
    # Mortgage receipts recorded by hand (no bank line) — the accountant may want to see these separately.
    manual = [t for t in transactions if t.bank_transaction_id is None and t.type not in ("advance", "funding")]

    return dict(
        year=year, month=month, start=start, end=end, label=start.strftime("%B %Y"),
        rows=rows, totals=totals, transactions=transactions, bank=bank, manual=manual,
        ytd=ytd_by_month(year, month), has_syndicated=any(r["share_pct"] != 100 for r in rows),
        generated=local_now(),
    )


def _status_at(m: Mortgage, as_of: date) -> str:
    if m.balance(as_of=as_of) <= 0 and m.funded_date <= as_of:
        return "Paid out"
    if m.maturity_date < as_of:
        return "Matured"
    return "Active"


def ytd_by_month(year: int, through_month: int):
    """Interest and fees received per mortgage per month, January → ``through_month``."""
    start = date(year, 1, 1)
    _, end = month_range(year, through_month)
    txns = MortgageTransaction.query.filter(MortgageTransaction.date.between(start, end)).all()
    table = {}
    for t in txns:
        if t.type in ("advance", "funding"):
            continue
        row = table.setdefault(t.mortgage_id, {"m": t.mortgage, "months": [ZERO] * through_month, "total": ZERO})
        amt = Decimal(str(t.interest or 0)) + Decimal(str(t.fees or 0))
        row["months"][t.date.month - 1] += amt
        row["total"] += amt
    rows = sorted(table.values(), key=lambda r: r["m"].reference)
    month_totals = [calc.money(sum((r["months"][i] for r in rows), ZERO)) for i in range(through_month)]
    return dict(
        labels=[date(year, i + 1, 1).strftime("%b") for i in range(through_month)],
        rows=rows, month_totals=month_totals, total=calc.money(sum(month_totals, ZERO)),
    )


# ----------------------------------------------------------------------------
# Exports
# ----------------------------------------------------------------------------
def transactions_csv(report: dict) -> str:
    out = io.StringIO()
    w = csv.writer(out)
    w.writerow(["Date", "Mortgage", "Borrower", "Property", "Type", "Amount", "Interest", "Principal", "Fees",
                "Bank account", "Bank description", "Notes"])
    for t in report["transactions"]:
        bt = t.bank_transaction
        w.writerow([t.date.isoformat()] + [csv_safe(v) for v in (
            t.mortgage.reference, t.mortgage.borrower_name, t.mortgage.property_address, t.type_label)] + [
            f"{t.signed_amount:.2f}", f"{Decimal(str(t.interest or 0)):.2f}",
            f"{Decimal(str(t.principal or 0)):.2f}", f"{Decimal(str(t.fees or 0)):.2f}"] + [csv_safe(v) for v in (
            bt.account_name if bt else "", bt.description if bt else "", t.notes or "")])
    return out.getvalue()


def month_end_workbook(report: dict, prepared_by: str = "") -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    MONEY = '#,##0.00;[Red]-#,##0.00'
    PCT = '0.00"%"'
    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="14213D")
    total_font = Font(bold=True)
    top_border = Border(top=Side(style="thin"), bottom=Side(style="double"))

    wb = Workbook()

    def sheet(title, heading, columns, first=False):
        ws = wb.active if first else wb.create_sheet()
        ws.title = title
        ws["A1"] = heading
        ws["A1"].font = Font(bold=True, size=14)
        ws["A2"] = (f"Period {report['start']:%b %d, %Y} – {report['end']:%b %d, %Y} · generated "
                    f"{report['generated']:%Y-%m-%d %H:%M}" + (f" by {prepared_by}" if prepared_by else ""))
        ws["A2"].font = Font(italic=True, color="666666")
        for i, (name, width, _) in enumerate(columns, start=1):
            c = ws.cell(row=4, column=i, value=name)
            c.font, c.fill = head_font, head_fill
            c.alignment = Alignment(wrap_text=True, vertical="center")
            ws.column_dimensions[get_column_letter(i)].width = width
        ws.row_dimensions[4].height = 30
        ws.freeze_panes = "A5"
        return ws

    def write_rows(ws, columns, rows, total_cols=()):
        r = 5
        for values in rows:
            for i, v in enumerate(values, start=1):
                c = set_text_cell(ws.cell(row=r, column=i), float(v) if isinstance(v, Decimal) else v)
                fmt = columns[i - 1][2]
                if fmt:
                    c.number_format = fmt
            r += 1
        if rows and total_cols:
            ws.cell(row=r, column=1, value="Total").font = total_font
            for i in total_cols:
                col = get_column_letter(i)
                c = ws.cell(row=r, column=i, value=f"=SUM({col}5:{col}{r - 1})")
                c.number_format, c.font, c.border = MONEY, total_font, top_border
        ws.auto_filter.ref = f"A4:{get_column_letter(len(columns))}{max(r - 1, 5)}"
        return r

    label = report["label"]
    t = report["totals"]

    # 1. Summary -----------------------------------------------------------------
    cols = [
        ("Reference", 11, None), ("Borrower", 26, None), ("Property", 30, None), ("Position", 8, None),
        ("Rate %", 8, PCT), ("Opening balance", 15, MONEY), ("Advances", 13, MONEY), ("Principal repaid", 14, MONEY),
        ("Closing balance", 15, MONEY), ("Interest received", 14, MONEY), ("Fees received", 12, MONEY),
        ("Total income (interest + fees)", 15, MONEY), ("Interest earned (accrual)", 15, MONEY),
        ("Scheduled payments due", 14, MONEY), ("Regular payments received", 14, MONEY),
        ("Variance", 12, MONEY), ("Arrears at month end", 13, MONEY), ("Status", 12, None),
    ]
    if report["has_syndicated"]:
        cols[12:12] = [("Your share %", 9, PCT), ("Your share of income", 14, MONEY)]
    ws = sheet("Summary", f"Mortgage income summary — {label}", cols, first=True)
    data = []
    for r in report["rows"]:
        m = r["m"]
        row = [m.reference, m.borrower_name, f"{m.property_address}, {m.property_city or ''}".strip(", "),
               m.position_label, r["rate"], r["opening"], r["advances"], r["principal"], r["closing"], r["interest"],
               r["fees"], r["income"], r["accrued"], r["expected"], r["regular"], r["variance"], r["arrears"], r["status"]]
        if report["has_syndicated"]:
            row[12:12] = [r["share_pct"], r["your_income"]]
        data.append(row)
    money_cols = [i for i, c in enumerate(cols, start=1) if c[2] == MONEY]
    end_row = write_rows(ws, cols, data, money_cols)
    # Headline numbers under the table.
    b = report["bank"]
    notes = [
        ("Interest received (cash)", t["interest"]),
        ("Fees received", t["fees"]),
        ("Total mortgage income received", t["income"]),
        ("Interest earned this month (accrual basis)", t["accrued"]),
        ("Principal repaid", t["principal"]),
        ("New advances", t["advances"]),
        ("Bank deposits this month", b["total"]),
        ("…matched to mortgages", b["matched"]),
        ("…marked not mortgage-related", b["ignored"]),
        ("…still unmatched (review)", b["unmatched"]),
    ]
    r0 = end_row + 2
    for i, (k, v) in enumerate(notes):
        ws.cell(row=r0 + i, column=2, value=k).font = Font(bold=i < 3)
        c = ws.cell(row=r0 + i, column=4, value=float(v))
        c.number_format = MONEY
    ws.cell(row=r0 + len(notes) + 1, column=2,
            value="Cash = money received in the month. Accrual = interest earned on each day's balance, "
                  "whether or not it was paid.").font = Font(italic=True, color="666666")

    # 2. Transactions --------------------------------------------------------------
    cols = [("Date", 11, "yyyy-mm-dd"), ("Reference", 11, None), ("Borrower", 26, None), ("Type", 18, None),
            ("Amount", 13, MONEY), ("Interest", 12, MONEY), ("Principal", 12, MONEY), ("Fees", 11, MONEY),
            ("Source", 9, None), ("Bank account", 18, None), ("Bank description", 40, None), ("Notes", 30, None)]
    ws = sheet("Transactions", f"Mortgage transactions — {label}", cols)
    data = []
    for tx in report["transactions"]:
        bt = tx.bank_transaction
        data.append([tx.date, tx.mortgage.reference, tx.mortgage.borrower_name, tx.type_label, tx.signed_amount,
                     Decimal(str(tx.interest or 0)), Decimal(str(tx.principal or 0)), Decimal(str(tx.fees or 0)),
                     "Bank" if bt else "Manual", bt.account_name if bt else "", bt.description if bt else "",
                     tx.notes or ""])
    write_rows(ws, cols, data, (5, 6, 7, 8))

    # 3. Bank reconciliation -------------------------------------------------------------
    cols = [("Date", 11, "yyyy-mm-dd"), ("Account", 18, None), ("Description", 45, None), ("Amount", 13, MONEY),
            ("Unallocated", 13, MONEY)]
    ws = sheet("Unmatched deposits", f"Deposits not yet matched to a mortgage — {label}", cols)
    data = [[bl.date, bl.account_name, bl.description, Decimal(str(bl.amount)), Decimal(str(bl.amount)) - bl.allocated]
            for bl in b["unmatched_lines"]]
    end_row = write_rows(ws, cols, data, (4, 5))
    if not data:
        ws.cell(row=5, column=1, value="None — every deposit this month is matched or marked as not mortgage-related.")

    # 4. Year to date ----------------------------------------------------------------
    ytd = report["ytd"]
    cols = [("Reference", 11, None), ("Borrower", 26, None)] + [(lbl, 12, MONEY) for lbl in ytd["labels"]] + [("Total", 14, MONEY)]
    ws = sheet("YTD income", f"Interest + fees received by month — {report['year']} year to date", cols)
    data = [[r["m"].reference, r["m"].borrower_name] + r["months"] + [r["total"]] for r in ytd["rows"]]
    write_rows(ws, cols, data, range(3, len(cols) + 1))

    for ws in wb.worksheets:
        ws.sheet_view.showGridLines = True
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToHeight = 0

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def previous_months(n=18, today: date | None = None):
    today = today or local_today()
    first = today.replace(day=1)
    return [(first - relativedelta(months=i)) for i in range(n)]
