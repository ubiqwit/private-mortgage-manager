"""Parse bank statement exports (CSV, OFX/QFX, XLSX) into normalised lines.

Canadian banks export CSVs in many shapes — with or without a header row, a single
signed amount column or separate debit/credit columns, MM/DD/YYYY or YYYY-MM-DD
dates. :func:`guess_mapping` makes a best guess that the user can correct in the UI
before anything is imported.
"""
from __future__ import annotations

import csv
import hashlib
import io
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation

DATE_FORMATS = [
    ("%Y-%m-%d", "YYYY-MM-DD"),
    ("%m/%d/%Y", "MM/DD/YYYY"),
    ("%d/%m/%Y", "DD/MM/YYYY"),
    ("%Y/%m/%d", "YYYY/MM/DD"),
    ("%m/%d/%y", "MM/DD/YY"),
    ("%d/%m/%y", "DD/MM/YY"),
    ("%m-%d-%Y", "MM-DD-YYYY"),
    ("%d-%m-%Y", "DD-MM-YYYY"),
    ("%Y%m%d", "YYYYMMDD"),
    ("%b %d, %Y", "Mon DD, YYYY"),
    ("%d %b %Y", "DD Mon YYYY"),
    ("%d-%b-%Y", "DD-Mon-YYYY"),
    ("%b %d %Y", "Mon DD YYYY"),
]

HEADER_WORDS = {
    "date": ["transaction date", "date posted", "posted date", "posting date", "trans date", "date"],
    "description": ["description 1", "description", "details", "transaction details", "payee", "memo", "name",
                    "transaction", "narrative", "particulars"],
    "description2": ["description 2", "memo", "additional info"],
    "amount": ["cad$", "amount", "transaction amount", "amount (cad)", "value"],
    "debit": ["debit", "withdrawal", "withdrawals", "paid out", "money out", "debits", "withdrawn"],
    "credit": ["credit", "deposit", "deposits", "paid in", "money in", "credits", "deposited"],
}


class StatementError(ValueError):
    pass


@dataclass
class ParsedLine:
    date: date
    description: str
    amount: Decimal
    fitid: str | None = None


@dataclass
class Mapping:
    date: int | None = None
    description: list[int] = field(default_factory=list)
    amount: int | None = None
    debit: int | None = None
    credit: int | None = None
    date_format: str | None = None
    has_header: bool = False
    header_row: int = 0
    flip_sign: bool = False

    def to_dict(self):
        return dict(self.__dict__)


# ----------------------------------------------------------------------------
# Low-level value parsing
# ----------------------------------------------------------------------------
def parse_amount(raw) -> Decimal | None:
    if raw is None:
        return None
    if isinstance(raw, (int, float, Decimal)):
        return Decimal(str(raw)).quantize(Decimal("0.01"))
    s = str(raw).strip()
    if not s:
        return None
    negative = False
    if s.startswith("(") and s.endswith(")"):
        negative, s = True, s[1:-1]
    up = s.upper()
    if up.endswith("CR"):
        s = s[:-2]
    elif up.endswith("DR"):
        negative, s = True, s[:-2]
    s = s.replace("$", "").replace(",", "").replace("CAD", "").replace(" ", "").strip()
    if s.endswith("-"):
        negative, s = True, s[:-1]
    if s.startswith("+"):
        s = s[1:]
    try:
        value = Decimal(s)
    except InvalidOperation:
        return None
    return (-value if negative else value).quantize(Decimal("0.01"))


def parse_date(raw, fmt: str | None = None) -> date | None:
    if raw is None:
        return None
    if isinstance(raw, datetime):
        return raw.date()
    if isinstance(raw, date):
        return raw
    s = str(raw).strip()
    if not s:
        return None
    s = re.sub(r"\s+", " ", s)
    formats = [fmt] if fmt else [f for f, _ in DATE_FORMATS]
    for f in formats:
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            continue
    # "2026-01-05 00:00:00" or "2026-01-05T00:00:00"
    m = re.match(r"^(\d{4}-\d{2}-\d{2})[ T]", s)
    if m:
        return datetime.strptime(m.group(1), "%Y-%m-%d").date()
    return None


def detect_date_format(values) -> str | None:
    """Pick the format that parses the most values (ties → the earlier, more common one).

    For slash dates both MM/DD and DD/MM may parse every value; a value with a day
    above 12 settles it, otherwise MM/DD (what Canadian bank exports use) wins.
    """
    values = [str(v).strip() for v in values if v not in (None, "") and str(v).strip()]
    if not values:
        return None
    best, best_count = None, 0
    for fmt, _ in DATE_FORMATS:
        count = sum(1 for v in values if parse_date(v, fmt))
        if count > best_count:
            best, best_count = fmt, count
    if best_count < max(1, int(len(values) * 0.8)):
        return None
    return best


# ----------------------------------------------------------------------------
# File readers
# ----------------------------------------------------------------------------
def decode(content: bytes) -> str:
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            return content.decode(enc)
        except UnicodeDecodeError:
            continue
    raise StatementError("Could not read the file's text encoding")


def file_kind(filename: str, content: bytes) -> str:
    name = (filename or "").lower()
    head = content[:2000].upper()
    if name.endswith((".ofx", ".qfx", ".qbo")) or b"<OFX>" in head or b"OFXHEADER" in head:
        return "ofx"
    if name.endswith((".xlsx", ".xlsm")) or content[:2] == b"PK":
        return "xlsx"
    if name.endswith(".pdf") or content[:4] == b"%PDF":
        return "pdf"
    return "csv"


def read_rows(filename: str, content: bytes) -> list[list[str]]:
    """Return the tabular rows of a CSV or XLSX file."""
    kind = file_kind(filename, content)
    if kind == "xlsx":
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        ws = wb.worksheets[0]
        rows = []
        for row in ws.iter_rows(values_only=True):
            cells = ["" if v is None else (v.strftime("%Y-%m-%d") if isinstance(v, (date, datetime)) else str(v))
                     for v in row]
            if any(c.strip() for c in cells):
                rows.append(cells)
        return rows
    if kind == "pdf":
        raise StatementError(
            "PDF statements can't be read reliably. In online banking, download the statement as CSV, "
            "Excel or OFX/QFX (Quicken/QuickBooks) instead — every major Canadian bank offers one of these."
        )
    if kind == "ofx":
        raise StatementError("OFX files are parsed directly, not as rows")
    text = decode(content)
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    rows = [r for r in csv.reader(io.StringIO(text), dialect) if any((c or "").strip() for c in r)]
    if not rows:
        raise StatementError("The file is empty")
    return rows


def parse_ofx(content: bytes) -> tuple[list[ParsedLine], str | None]:
    """Parse OFX 1.x (SGML) or 2.x (XML) statements. Returns lines and the account id."""
    text = decode(content)
    acct = re.search(r"<ACCTID>([^<\r\n]+)", text, re.I)
    lines = []
    for block in re.findall(r"<STMTTRN>(.*?)(?:</STMTTRN>|(?=<STMTTRN>)|(?=</BANKTRANLIST>))", text, re.S | re.I):
        def tag(name, block=block):
            m = re.search(rf"<{name}>([^<\r\n]*)", block, re.I)
            return m.group(1).strip() if m else ""

        d = parse_date(tag("DTPOSTED")[:8], "%Y%m%d")
        amt = parse_amount(tag("TRNAMT"))
        if d is None or amt is None:
            continue
        desc = " ".join(x for x in (tag("NAME"), tag("MEMO")) if x) or tag("TRNTYPE")
        lines.append(ParsedLine(date=d, description=desc, amount=amt, fitid=tag("FITID") or None))
    if not lines:
        raise StatementError("No transactions found in the OFX/QFX file")
    return lines, (acct.group(1).strip() if acct else None)


# ----------------------------------------------------------------------------
# Column mapping
# ----------------------------------------------------------------------------
def _norm(s) -> str:
    return re.sub(r"\s+", " ", str(s or "").strip().lower())


def _header_index(rows):
    """Index of the header row within the first 15 rows, or None if the file has none."""
    for i, row in enumerate(rows[:15]):
        cells = [_norm(c) for c in row]
        hits = 0
        for c in cells:
            if any(c == w or c.startswith(w) for words in HEADER_WORDS.values() for w in words):
                hits += 1
        has_date = any("date" in c for c in cells)
        looks_numeric = any(re.match(r"^[-+$(]?\d", c) for c in cells if c)
        if hits >= 2 and has_date and not looks_numeric:
            return i
    return None


def _find(cells, role, exclude=()):
    for word in HEADER_WORDS[role]:
        for i, c in enumerate(cells):
            if i in exclude:
                continue
            if c == word:
                return i
    for word in HEADER_WORDS[role]:
        for i, c in enumerate(cells):
            if i in exclude:
                continue
            if c.startswith(word) or (len(word) > 4 and word in c):
                return i
    return None


def guess_mapping(rows: list[list[str]]) -> Mapping:
    m = Mapping()
    header = _header_index(rows)
    width = max(len(r) for r in rows)
    if header is not None:
        m.has_header, m.header_row = True, header
        cells = [_norm(c) for c in rows[header]]
        m.date = _find(cells, "date")
        used = {m.date}
        m.amount = _find(cells, "amount", used)
        if m.amount is not None:
            used.add(m.amount)
        else:
            m.debit = _find(cells, "debit", used)
            used.add(m.debit)
            m.credit = _find(cells, "credit", used)
            used.add(m.credit)
        desc = _find(cells, "description", used)
        if desc is not None:
            m.description.append(desc)
            used.add(desc)
            d2 = _find(cells, "description2", used)
            if d2 is not None:
                m.description.append(d2)
        data = rows[header + 1:]
    else:
        data = rows
        cols = list(range(width))

        def column(i):
            return [r[i] if i < len(r) else "" for r in data[:50]]

        date_cols = [i for i in cols if detect_date_format(column(i))]
        m.date = date_cols[0] if date_cols else None
        numeric = []
        for i in cols:
            if i == m.date:
                continue
            vals = [v for v in column(i) if str(v).strip()]
            if vals and sum(1 for v in vals if parse_amount(v) is not None) >= 0.8 * len(vals):
                numeric.append(i)
        text_cols = [i for i in cols if i != m.date and i not in numeric]
        # Keep every text column that carries real words (skip "-" fillers and codes), max two.
        wordy = [i for i in text_cols
                 if sum(len(str(v).strip()) for v in column(i)) / max(len(column(i)), 1) >= 3]
        m.description = (wordy or text_cols)[:2]
        # Sparse numeric columns (blank on some rows) are debit/credit pairs; dense ones are amount/balance.
        sparse = [i for i in numeric if any(not str(v).strip() for v in column(i))]
        if len(sparse) >= 2:
            m.debit, m.credit = sparse[0], sparse[1]
        elif numeric:
            m.amount = numeric[0]
    if m.date is not None:
        m.date_format = detect_date_format([r[m.date] for r in data[:200] if m.date < len(r)])
    return m


def apply_mapping(rows: list[list[str]], m: Mapping) -> tuple[list[ParsedLine], list[str]]:
    """Convert rows to lines. Returns (lines, warnings for skipped rows)."""
    if m.date is None:
        raise StatementError("Choose which column holds the date")
    if m.amount is None and m.debit is None and m.credit is None:
        raise StatementError("Choose an amount column, or debit and credit columns")
    data = rows[m.header_row + 1:] if m.has_header else rows
    lines, warnings = [], []
    for n, row in enumerate(data, start=(m.header_row + 2 if m.has_header else 1)):
        def cell(i, row=row):
            return row[i] if i is not None and i < len(row) else ""

        d = parse_date(cell(m.date), m.date_format)
        if d is None:
            if any(str(c).strip() for c in row):
                warnings.append(f"Row {n}: skipped (no valid date in '{cell(m.date)}')")
            continue
        if m.amount is not None:
            amount = parse_amount(cell(m.amount))
        else:
            debit = parse_amount(cell(m.debit)) or Decimal(0)
            credit = parse_amount(cell(m.credit)) or Decimal(0)
            amount = (abs(credit) - abs(debit)) if (debit or credit) else None
        if amount is None:
            warnings.append(f"Row {n}: skipped (no amount)")
            continue
        if m.flip_sign:
            amount = -amount
        desc = " ".join(str(cell(i)).strip() for i in m.description if str(cell(i)).strip())
        lines.append(ParsedLine(date=d, description=re.sub(r"\s+", " ", desc), amount=amount.quantize(Decimal("0.01"))))
    if not lines:
        raise StatementError("No transactions could be read with this column mapping")
    return lines, warnings


def fingerprints(lines: list[ParsedLine], account: str) -> list[str]:
    """Stable ids used to skip lines that were already imported.

    Identical lines on the same day (two equal e-transfers) are kept apart by their
    occurrence number within the file.
    """
    seen: dict[str, int] = {}
    out = []
    for ln in lines:
        if ln.fitid:
            base = f"fitid|{account}|{ln.fitid}"
        else:
            base = f"{account}|{ln.date.isoformat()}|{ln.amount}|{_norm(ln.description)}"
        seen[base] = seen.get(base, 0) + 1
        out.append(hashlib.sha256(f"{base}|{seen[base]}".encode()).hexdigest())
    return out
