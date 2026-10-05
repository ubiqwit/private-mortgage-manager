"""Full data export: every table the owner cares about, in one Excel workbook."""
from __future__ import annotations

import io
from datetime import date, datetime
from decimal import Decimal

from ..models import (
    AuditLog,
    BankTransaction,
    Company,
    MarketObservation,
    Mortgage,
    MortgageActivity,
    MortgageTransaction,
    StatementImport,
    TermHistory,
    User,
)
from ..timeutil import now as local_now
from .safety import set_text_cell

SHEETS = [
    ("Companies", Company, None),
    ("Mortgages", Mortgage, None),
    ("Transactions", MortgageTransaction, None),
    ("Term history", TermHistory, None),
    ("Activity log", MortgageActivity, None),
    ("Bank lines", BankTransaction, None),
    ("Statement imports", StatementImport, None),
    ("Market rates", MarketObservation, None),
    ("Users", User, {"password_hash", "password_stamp"}),
    ("Audit log", AuditLog, None),
]


def _cell(value):
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, (datetime, date)) or value is None or isinstance(value, (int, float, str, bool)):
        return value
    return str(value)


def export_workbook() -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    wb.remove(wb.active)
    info = wb.create_sheet("About")
    info["A1"] = "Private Mortgage Manager — full data export"
    info["A1"].font = Font(bold=True, size=14)
    info["A2"] = (f"Exported {local_now():%Y-%m-%d %H:%M}. Every company, one sheet per table "
                "(company_id links rows to the Companies sheet); keep this file somewhere safe.")
    for name, model, exclude in SHEETS:
        ws = wb.create_sheet(name)
        cols = [c.name for c in model.__table__.columns if not exclude or c.name not in exclude]
        ws.append(cols)
        for c in ws[1]:
            c.font, c.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="14213D")
        rows = model.query.order_by(*model.__table__.primary_key.columns).execution_options(skip_company_scope=True)
        for r, row in enumerate(rows.all(), start=2):
            for i, c in enumerate(cols, start=1):
                set_text_cell(ws.cell(row=r, column=i), _cell(getattr(row, c)))
        for i, col in enumerate(cols, start=1):
            ws.column_dimensions[get_column_letter(i)].width = max(10, min(40, len(col) + 4))
        ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
