"""Each mortgage records who owns it; the list filters by owner and reports show it."""
import io
from datetime import date

import sqlalchemy as sa
from flask_migrate import downgrade, upgrade
from openpyxl import load_workbook

from app import MIGRATIONS_DIR, create_app, db
from app.models import Mortgage, Setting
from app.services import importer, portfolio, reports

from .test_import import ROW as IMPORT_ROW
from .test_import import csv_with, post
from .test_mortgages import create


def test_owners_saved_shown_and_filtered(client):
    a = create(client, owners="9929916 Canada Inc", borrower_name="Alpha Borrower")
    create(client, owners="Sal & Suresh 50-50", borrower_name="Beta Borrower")
    create(client, borrower_name="Gamma Borrower")
    assert a.owners == "9929916 Canada Inc"
    assert b"9929916 Canada Inc" in client.get(f"/mortgages/{a.id}").data

    page = client.get("/mortgages/").get_data(as_text=True)
    assert "All owners" in page and "Sal &amp; Suresh 50-50" in page
    only = client.get("/mortgages/?owner=9929916 Canada Inc").get_data(as_text=True)
    assert "Alpha Borrower" in only and "Beta Borrower" not in only and "Gamma Borrower" not in only
    found = client.get("/mortgages/?q=suresh").get_data(as_text=True)
    assert "Beta Borrower" in found and "Alpha Borrower" not in found

    # The form suggests owners already in use.
    assert 'list="dl_owners"' in client.get("/mortgages/new").get_data(as_text=True)


def test_owner_cleared_when_blank(client):
    m = create(client, owners="Somebody")
    form = client.get(f"/mortgages/{m.id}/edit")
    assert form.status_code == 200
    resp = client.post(f"/mortgages/{m.id}/edit", data={"owners": "", "borrower_name": m.borrower_name,
                                                        "property_address": m.property_address,
                                                        "principal_amount": "200000", "interest_rate": "10",
                                                        "funded_date": "2026-01-01", "term_months": "12"})
    assert resp.status_code == 302
    assert db.session.get(Mortgage, m.id).owners in (None, "")


def test_dashboard_concentration_by_owner(client):
    create(client)
    assert portfolio.dashboard_stats()["by_owner"] == []  # nothing recorded yet: no owners panel
    create(client, owners="Suresh Malhotra")
    owners = dict(portfolio.dashboard_stats()["by_owner"])
    assert set(owners) == {"Suresh Malhotra", "Not recorded"}
    assert "Owners" in client.get("/").get_data(as_text=True)


def test_reports_include_owners(client):
    create(client, owners="Suresh Malhotra", funded_date="2026-01-01")
    r = reports.month_end_report(2026, 3)
    ws = load_workbook(io.BytesIO(reports.month_end_workbook(r)))["Summary"]
    rows = list(ws.iter_rows(values_only=True))
    header = next(row for row in rows if row and row[0] == "Reference")
    data = rows[rows.index(header) + 1]
    assert header[2] == "Owners" and data[2] == "Suresh Malhotra"
    annual = load_workbook(io.BytesIO(reports.annual_workbook(reports.annual_report(2026)))).active
    assert annual.cell(row=4, column=3).value == "Owners"
    assert annual.cell(row=5, column=3).value == "Suresh Malhotra"
    assert "Suresh Malhotra" in client.get("/reports/month-end?month=2026-03").get_data(as_text=True)


def test_import_owners_column(client):
    resp = post(client, csv_with({**IMPORT_ROW, "Owners": "9929916 Canada Inc"}))
    assert resp.status_code == 302
    assert Mortgage.query.one().owners == "9929916 Canada Inc"


def test_template_example_row_matches_columns(client):
    ws = load_workbook(io.BytesIO(importer.template_workbook()))["Mortgages"]
    rows = list(ws.iter_rows(values_only=True))
    assert len(rows[0]) == len(importer.COLUMNS)
    assert dict(zip(rows[0], rows[1], strict=True))["Owners"] == "Example Holdings Inc"


def test_upgrade_from_companies_keeps_company_as_owner(tmp_path):
    """A database that used companies (the previous release) moves each company's name
    into its mortgages' owners, and renames references that were only unique per company."""
    url = f"sqlite:///{tmp_path / 'companies.db'}"
    app = create_app({"SECRET_KEY": "t", "SQLALCHEMY_DATABASE_URI": url,
                      "MARKET_FETCH_ENABLED": False})
    with app.app_context():
        for ref in ("M-001", "M-002", "M-003"):
            db.session.add(Mortgage(reference=ref, borrower_name=ref, property_address="1 Main St",
                                    principal_amount=100000, interest_rate=10, funded_date=date(2026, 1, 1),
                                    first_payment_date=date(2026, 2, 1), maturity_date=date(2027, 1, 1)))
        db.session.commit()
        db.session.remove()

        downgrade(directory=MIGRATIONS_DIR, revision="b376b2c09d2e")  # back to the companies release
        with db.engine.begin() as conn:
            assert "company" in sa.inspect(conn).get_table_names()
            conn.execute(sa.text("UPDATE company SET name = 'MFT Holdings Inc.' WHERE id = 1"))
            conn.execute(sa.text("INSERT INTO company (id, name, created_at, updated_at) "
                                 "VALUES (2, 'Suresh Malhotra', '2026-10-05', '2026-10-05')"))
            # Suresh's company had its own M-001.
            conn.execute(sa.text("UPDATE mortgage SET company_id = 2, reference = 'M-001' WHERE reference = 'M-003'"))
            conn.execute(sa.text("DELETE FROM setting WHERE key LIKE 'books_closed_through%'"))
            conn.execute(sa.text("INSERT INTO setting (key, value) VALUES ('books_closed_through:1', '2026-08-31'), "
                                 "('books_closed_through:2', '2026-06-30')"))

        upgrade(directory=MIGRATIONS_DIR)
        db.session.remove()
        rows = {m.borrower_name: m for m in Mortgage.query.all()}
        assert rows["M-001"].owners == "MFT Holdings Inc." and rows["M-001"].reference == "M-001"
        assert rows["M-002"].owners == "MFT Holdings Inc."
        assert rows["M-003"].owners == "Suresh Malhotra" and rows["M-003"].reference == "M-001-2"
        assert Setting.get("books_closed_through") == "2026-08-31"
        assert "company" not in sa.inspect(db.engine).get_table_names()


def test_upgrade_from_single_default_company_leaves_owners_blank(tmp_path):
    url = f"sqlite:///{tmp_path / 'one.db'}"
    app = create_app({"SECRET_KEY": "t", "SQLALCHEMY_DATABASE_URI": url,
                      "MARKET_FETCH_ENABLED": False})
    with app.app_context():
        db.session.add(Mortgage(reference="M-001", borrower_name="X", property_address="1 Main St",
                                principal_amount=100000, interest_rate=10, funded_date=date(2026, 1, 1),
                                first_payment_date=date(2026, 2, 1), maturity_date=date(2027, 1, 1)))
        db.session.commit()
        db.session.remove()
        downgrade(directory=MIGRATIONS_DIR, revision="b376b2c09d2e")
        upgrade(directory=MIGRATIONS_DIR)
        db.session.remove()
        assert Mortgage.query.one().owners is None
