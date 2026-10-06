"""Mortgagees: who owns each mortgage and their percentage."""
import io
import json
from datetime import date
from decimal import Decimal

import pytest
import sqlalchemy as sa
from flask_migrate import downgrade, upgrade
from openpyxl import load_workbook
from werkzeug.datastructures import MultiDict

from app import MIGRATIONS_DIR, create_app, db
from app.models import Mortgage, Mortgagee, Setting, User
from app.services import importer, portfolio, reports
from app.services import mortgagees as mg

from .test_import import ROW as IMPORT_ROW
from .test_import import csv_with, post
from .test_mortgages import BASE, create


def owners(m):
    return [(o.name, o.share_pct) for o in db.session.get(Mortgage, m.id).mortgagees]


def with_mortgagees(*rows):
    data = MultiDict(BASE)
    data.add("mortgagees_form", "1")
    for name, pct in rows:
        data.add("mortgagee_name", name)
        data.add("mortgagee_pct", pct)
    return data


# --- parsing & validation ---------------------------------------------------------

@pytest.mark.parametrize("text, expected", [
    ("", []),
    ("Suresh Malhotra", [("Suresh Malhotra", 100)]),
    ("9929916 Canada Inc 50%; Suresh Malhotra 50%", [("9929916 Canada Inc", 50), ("Suresh Malhotra", 50)]),
    ("1022409 50%, CFT 33.7%, 9929916 16.3%", [("1022409", 50), ("CFT", Decimal("33.7")), ("9929916", Decimal("16.3"))]),
    ("Sal; Suresh", [("Sal", 50), ("Suresh", 50)]),
    ("A 40%; B; C", [("A", 40), ("B", 30), ("C", 30)]),
    ("1022409 Ontario Inc. 50%", [("1022409 Ontario Inc.", 50)]),
    ("Smith, Jones & Co", [("Smith, Jones & Co", 100)]),  # a comma inside a name isn't a separator
])
def test_parse_text(text, expected):
    assert mg.parse_text(text) == [(n, Decimal(p)) for n, p in expected]


@pytest.mark.parametrize("pairs, message", [
    ([("A", "60"), ("B", "50")], "can't exceed 100%"),
    ([("A", "100"), ("B", "")], "already add up to 100%"),
    ([("A", "50"), ("a", "50")], "listed twice"),
    ([("", "50")], "needs a name"),
    ([("A", "0")], "more than 0%"),
    ([("A", "abc")], "not a number"),
])
def test_validation(pairs, message):
    with pytest.raises(mg.MortgageeError, match=message):
        mg.complete(pairs)


def test_blank_shares_split_the_rest_exactly():
    pairs = mg.complete([("A", ""), ("B", ""), ("C", "")])
    assert sum(p for _, p in pairs) == 100 and pairs[0][1] == Decimal("33.33")


# --- form --------------------------------------------------------------------------

def test_create_and_edit_mortgagees(client):
    resp = client.post("/mortgages/new", data=with_mortgagees(("9929916 Canada Inc", "60"), ("Suresh Malhotra", ""),
                                                              ("", "")))
    assert resp.status_code == 302
    m = Mortgage.query.one()
    assert owners(m) == [("9929916 Canada Inc", 60), ("Suresh Malhotra", 40)]

    form = client.get(f"/mortgages/{m.id}/edit").get_data(as_text=True)
    assert 'value="9929916 Canada Inc"' in form and 'value="60"' in form
    assert '<option value="Suresh Malhotra">' in form  # suggestions

    resp = client.post(f"/mortgages/{m.id}/edit", data=with_mortgagees(("Suresh Malhotra", "100")))
    assert resp.status_code == 302
    assert owners(m) == [("Suresh Malhotra", 100)]
    assert Mortgagee.query.count() == 1  # replaced, not added

    client.post(f"/mortgages/{m.id}/edit", data=with_mortgagees(("", "")))
    assert owners(m) == []


def test_invalid_shares_are_rejected_and_kept_on_the_form(client):
    m = create(client)
    resp = client.post(f"/mortgages/{m.id}/edit", data=with_mortgagees(("A", "70"), ("B", "40")))
    assert resp.status_code == 400
    page = resp.get_data(as_text=True)
    assert "can&#39;t exceed 100%" in page or "can't exceed 100%" in page
    assert 'value="B"' in page and 'value="40"' in page  # what they typed is still there
    assert owners(m) == []


def test_hidden_section_keeps_mortgagees(client):
    client.post("/mortgages/new", data=with_mortgagees(("A", "100")))
    m = Mortgage.query.one()
    client.post(f"/mortgages/{m.id}/edit", data=BASE)  # a form without the section
    assert owners(m) == [("A", 100)]


def test_old_owners_form_setting_still_shows_mortgagees(client):
    user = User.query.one()
    user.form_fields = json.dumps(["owners", "notes"])
    db.session.commit()
    assert 'name="mortgagees_form"' in client.get("/mortgages/new").get_data(as_text=True)


# --- list, detail, dashboard, reports, import --------------------------------------

def test_list_columns_filter_and_share(client):
    client.post("/mortgages/new", data=with_mortgagees(("9929916 Canada Inc", "50"), ("Suresh Malhotra", "50")))
    client.post("/mortgages/new", data={**with_mortgagees(("Suresh Malhotra", "25")), "property_address": "9 Other Rd"})
    client.post("/mortgages/new", data={**BASE, "property_address": "1 Nobody Lane"})

    page = client.get("/mortgages/").get_data(as_text=True)
    head = page[page.index("<thead"):page.index("</thead>")]
    for gone in (">Ref<", ">Borrower<", ">Arrears<", ">Next due<"):
        assert gone not in head
    assert "Original principal" in head and "Balance" in head and "Mortgagees" in head
    assert "Jane Smith" not in page  # borrower not shown

    only = client.get("/mortgages/?mortgagee=Suresh Malhotra").get_data(as_text=True)
    assert "12 Maple Ave" in only and "9 Other Rd" in only and "1 Nobody Lane" not in only
    # 50% + 25% of two $200,000 loans.
    assert "Suresh Malhotra&#39;s share of the balance" in only or "Suresh Malhotra's share of the balance" in only
    assert "$150,000.00" in only
    assert "9 Other Rd" in client.get("/mortgages/?q=suresh").get_data(as_text=True)


def test_detail_shows_shares_of_balance(client):
    client.post("/mortgages/new", data=with_mortgagees(("A", "75")))
    m = Mortgage.query.one()
    page = client.get(f"/mortgages/{m.id}").get_data(as_text=True)
    assert "Mortgagees" in page and "$150,000.00" in page  # 75% of 200,000
    assert "Not assigned" in page and "$50,000.00" in page


def test_dashboard_splits_balance_by_share(client):
    create(client)
    assert portfolio.dashboard_stats()["by_owner"] == []
    client.post("/mortgages/new", data=with_mortgagees(("A", "60"), ("B", "")))
    split = dict(portfolio.dashboard_stats()["by_owner"])
    assert split == {"A": Decimal("120000.00"), "B": Decimal("80000.00"), "Not recorded": Decimal("200000.00")}


def test_reports_list_mortgagees(client):
    client.post("/mortgages/new", data=with_mortgagees(("A", "60"), ("B", "40")))
    ws = load_workbook(io.BytesIO(reports.month_end_workbook(reports.month_end_report(2026, 3))))["Summary"]
    rows = list(ws.iter_rows(values_only=True))
    header = next(row for row in rows if row and row[0] == "Reference")
    assert header[2] == "Mortgagees" and rows[rows.index(header) + 1][2] == "A 60%, B 40%"
    annual = load_workbook(io.BytesIO(reports.annual_workbook(reports.annual_report(2026)))).active
    assert annual.cell(row=4, column=3).value == "Mortgagees" and annual.cell(row=5, column=3).value == "A 60%, B 40%"


def test_import_mortgagees(client):
    resp = post(client, csv_with({**IMPORT_ROW, "Mortgagees": "9929916 Canada Inc 50%; Suresh Malhotra 50%"}))
    assert resp.status_code == 302
    assert owners(Mortgage.query.one()) == [("9929916 Canada Inc", 50), ("Suresh Malhotra", 50)]


def test_import_older_owners_header(client):
    content = (b"Borrower name,Property address,Principal,Rate %,Funded date,Term months,Owners\n"
               b"Jane,12 Maple Ave,250000,10,2026-01-01,12,9929916 Canada Inc\n")
    assert post(client, content).status_code == 302
    assert owners(Mortgage.query.one()) == [("9929916 Canada Inc", 100)]


def test_import_rejects_bad_shares(client):
    resp = post(client, csv_with({**IMPORT_ROW, "Mortgagees": "A 80%; B 80%"}))
    assert resp.status_code == 400 and Mortgage.query.count() == 0


def test_template_example_row_matches_columns(client):
    ws = load_workbook(io.BytesIO(importer.template_workbook()))["Mortgages"]
    rows = list(ws.iter_rows(values_only=True))
    assert len(rows[0]) == len(importer.COLUMNS)
    assert dict(zip(rows[0], rows[1], strict=True))["Mortgagees"] == "Example Holdings Inc 60%; Jane Lender 40%"


# --- upgrades ------------------------------------------------------------------------

def _migrated_app(tmp_path, name):
    return create_app({"SECRET_KEY": "t", "SQLALCHEMY_DATABASE_URI": f"sqlite:///{tmp_path / name}",
                       "MARKET_FETCH_ENABLED": False})


def _add_mortgages(*refs):
    for ref in refs:
        db.session.add(Mortgage(reference=ref, borrower_name=ref, property_address="1 Main St",
                                principal_amount=100000, interest_rate=10, funded_date=date(2026, 1, 1),
                                first_payment_date=date(2026, 2, 1), maturity_date=date(2027, 1, 1)))
    db.session.commit()
    db.session.remove()


def test_upgrade_from_companies_makes_each_company_a_mortgagee(tmp_path):
    app = _migrated_app(tmp_path, "companies.db")
    with app.app_context():
        _add_mortgages("M-001", "M-002", "M-003")
        downgrade(directory=MIGRATIONS_DIR, revision="b376b2c09d2e")  # the companies release
        with db.engine.begin() as conn:
            conn.execute(sa.text("UPDATE company SET name = 'MFT Holdings Inc.' WHERE id = 1"))
            conn.execute(sa.text("INSERT INTO company (id, name, created_at, updated_at) "
                                 "VALUES (2, 'Suresh Malhotra', '2026-10-05', '2026-10-05')"))
            conn.execute(sa.text("UPDATE mortgage SET company_id = 2, reference = 'M-001' WHERE reference = 'M-003'"))
            conn.execute(sa.text("DELETE FROM setting WHERE key LIKE 'books_closed_through%'"))
            conn.execute(sa.text("INSERT INTO setting (key, value) VALUES ('books_closed_through:1', '2026-08-31'), "
                                 "('books_closed_through:2', '2026-06-30')"))
        upgrade(directory=MIGRATIONS_DIR)
        db.session.remove()
        rows = {m.borrower_name: m for m in Mortgage.query.all()}
        assert owners(rows["M-001"]) == [("MFT Holdings Inc.", 100)]
        assert owners(rows["M-003"]) == [("Suresh Malhotra", 100)] and rows["M-003"].reference == "M-001-2"
        assert Setting.get("books_closed_through") == "2026-08-31"
        assert "company" not in sa.inspect(db.engine).get_table_names()


def test_upgrade_from_owners_text(tmp_path):
    app = _migrated_app(tmp_path, "owners.db")
    with app.app_context():
        _add_mortgages("M-001", "M-002")
        downgrade(directory=MIGRATIONS_DIR, revision="77dff3b1fd16")  # the owners-text release
        with db.engine.begin() as conn:
            conn.execute(sa.text("UPDATE mortgage SET owners = 'Sal & Suresh 50-50' WHERE reference = 'M-001'"))
        upgrade(directory=MIGRATIONS_DIR)
        db.session.remove()
        rows = {m.reference: m for m in Mortgage.query.all()}
        assert owners(rows["M-001"]) == [("Sal & Suresh 50-50", 100)]
        assert owners(rows["M-002"]) == []
