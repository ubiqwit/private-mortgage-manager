import io
from datetime import date
from decimal import Decimal

from openpyxl import load_workbook

from app import db
from app.models import Mortgage
from app.services import reports

from .test_mortgages import create


def setup_book(client):
    m = create(client)  # $200k @ 10% monthly IO, funded Jan 1 2026, $4,000 lender fee
    for d in ("2026-02-01", "2026-03-01"):
        client.post(f"/mortgages/{m.id}/transactions", data={"type": "payment", "date": d, "amount": "1666.67"})
    client.post(f"/mortgages/{m.id}/transactions", data={"type": "prepayment", "date": "2026-03-16", "amount": "50000"})
    return db.session.get(Mortgage, m.id)


def test_january_includes_funding_and_partial_accrual(client):
    setup_book(client)
    r = reports.month_end_report(2026, 1)
    row = r["rows"][0]
    assert row["opening"] == 0 and row["advances"] == Decimal("200000.00") and row["closing"] == Decimal("200000.00")
    assert row["fees"] == Decimal("4000.00") and row["income"] == Decimal("4000.00")
    # 30 of 31 days earn interest (none on the funding day).
    assert row["accrued"] == Decimal("1612.90")


def test_february_full_month(client):
    setup_book(client)
    r = reports.month_end_report(2026, 2)
    row = r["rows"][0]
    assert row["interest"] == Decimal("1666.67")
    assert row["accrued"] == Decimal("1666.67")
    assert row["expected"] == row["regular"] == Decimal("1666.67")
    assert r["totals"]["income"] == Decimal("1666.67")


def test_march_prepayment_reduces_accrual(client):
    setup_book(client)
    r = reports.month_end_report(2026, 3)
    row = r["rows"][0]
    assert row["principal"] == Decimal("50000.00")
    assert row["closing"] == Decimal("150000.00")
    # Mar 1–16 on $200k (the receipt day still earns on the old balance), Mar 17–31 on $150k.
    expected = (Decimal(200000) * 16 + Decimal(150000) * 15) * Decimal("0.10") / 12 / 31
    assert row["accrued"] == expected.quantize(Decimal("0.01"))


def test_ytd_and_exports(client):
    setup_book(client)
    r = reports.month_end_report(2026, 3)
    assert r["ytd"]["labels"] == ["Jan", "Feb", "Mar"]
    assert r["ytd"]["total"] == Decimal("4000.00") + Decimal("1666.67") * 2

    wb = load_workbook(io.BytesIO(reports.month_end_workbook(r, "Tester")))
    assert wb.sheetnames == ["Summary", "Transactions", "Unmatched deposits", "YTD income"]
    ws = wb["Summary"]
    assert ws["A5"].value == "M-001"
    assert ws["A6"].value == "Total" and str(ws["J6"].value).startswith("=SUM(")


def test_report_pages_and_downloads(client):
    setup_book(client)
    assert client.get("/reports/month-end?month=2026-03").status_code == 200
    assert client.get("/reports/month-end").status_code == 200
    resp = client.get("/reports/month-end.xlsx?month=2026-03")
    assert resp.status_code == 200 and resp.data[:2] == b"PK"
    assert "mortgage-income-2026-03.xlsx" in resp.headers["Content-Disposition"]
    assert client.get("/reports/month-end?month=2026-13").status_code == 400
    page = client.get("/reports/month-end?month=2026-03").get_data(as_text=True)
    assert "Annual summary" not in page and "CSV" not in page and "Variance" not in page
    assert client.get("/reports/month-end-transactions.csv?month=2026-03").status_code == 404
    assert client.get("/reports/annual?year=2026").status_code == 404


def test_default_month_is_previous():
    assert reports.default_month(date(2026, 10, 2)) == (2026, 9)
    assert reports.default_month(date(2026, 1, 15)) == (2025, 12)


def test_payout_month_has_no_payment_due_after_payout(client):
    m = setup_book(client)
    client.post(f"/mortgages/{m.id}/transactions", data={"type": "payment", "date": "2026-04-01", "amount": "1250"})
    client.post(f"/mortgages/{m.id}/transactions", data={"type": "payout", "date": "2026-05-01", "amount": "151250"})
    r = reports.month_end_report(2026, 5)
    row = r["rows"][0]
    assert row["closing"] == 0 and row["expected"] == 0 and row["interest"] == Decimal("1250.00")
    assert reports.month_end_report(2026, 6)["rows"] == []


def test_mortgagee_share_report(client):
    from werkzeug.datastructures import MultiDict

    from .test_mortgages import BASE

    def new(address, *owners):
        data = MultiDict({**BASE, "property_address": address, "record_lender_fee": ""})
        data.add("mortgagees_form", "1")
        for name, pct in owners:
            data.add("mortgagee_name", name)
            data.add("mortgagee_pct", pct)
        assert client.post("/mortgages/new", data=data).status_code == 302
        m = Mortgage.query.filter_by(property_address=address).one()
        client.post(f"/mortgages/{m.id}/transactions", data={"type": "payment", "date": "2026-02-01", "amount": "1666.67"})
        return m

    new("1 Shared St", ("9929916 Canada Inc", "60"), ("Suresh Malhotra", "40"))
    new("2 Solo Ave", ("Suresh Malhotra", "100"))
    new("3 Other Rd", ("1022409 Ontario Inc.", "100"))

    whole = reports.month_end_report(2026, 2)
    assert len(whole["rows"]) == 3 and whole["totals"]["interest"] == Decimal("1666.67") * 3
    assert whole["bank"] is not None

    r = reports.month_end_report(2026, 2, mortgagee="Suresh Malhotra")
    rows = {row["m"].property_address: row for row in r["rows"]}
    assert set(rows) == {"1 Shared St", "2 Solo Ave"}
    assert rows["1 Shared St"]["share_pct"] == Decimal("40.00")
    assert rows["1 Shared St"]["interest"] == Decimal("666.67")      # 40% of 1,666.67
    assert rows["1 Shared St"]["closing"] == Decimal("80000.00")     # 40% of 200,000
    assert rows["2 Solo Ave"]["interest"] == Decimal("1666.67")
    assert r["totals"]["interest"] == Decimal("2333.34")
    assert r["bank"] is None and len(r["transactions"]) == 2         # their mortgages only, full amounts
    assert r["ytd"]["total"] == Decimal("2333.34")

    page = client.get("/reports/month-end?month=2026-02&mortgagee=Suresh Malhotra").get_data(as_text=True)
    assert "Suresh Malhotra&#39;s share" in page and ">M-003<" not in page and "Bank reconciliation" not in page
    assert "Close books" not in page  # closing applies to the whole book, so it's offered there
    resp = client.get("/reports/month-end.xlsx?month=2026-02&mortgagee=Suresh Malhotra")
    assert "mortgage-income-2026-02-suresh-malhotra.xlsx" in resp.headers["Content-Disposition"]
    wb = load_workbook(io.BytesIO(resp.data))
    assert wb.sheetnames == ["Summary", "Transactions", "YTD income"]
    ws = wb["Summary"]
    assert ws["A1"].value == "Mortgage income summary — February 2026 — Suresh Malhotra's share"
    header = [c.value for c in ws[4]]
    assert header[6] == "Share %" and "Variance" not in header
    # An unknown name falls back to the whole book.
    assert ">M-003<" in client.get("/reports/month-end?month=2026-02&mortgagee=Nobody").get_data(as_text=True)
