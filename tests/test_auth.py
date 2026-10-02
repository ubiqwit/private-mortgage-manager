from app.models import User

from .conftest import PASSWORD, make_user


def test_login_required(anon_client):
    resp = anon_client.get("/")
    assert resp.status_code == 302 and "/login" in resp.headers["Location"]
    assert anon_client.get("/healthz").status_code == 200


def test_login_and_logout(client):
    assert client.get("/").status_code == 200
    client.post("/logout")
    assert client.get("/").status_code == 302


def test_bad_password_and_lockout(app, anon_client):
    make_user()
    for _ in range(5):
        r = anon_client.post("/login", data={"email": "owner@example.com", "password": "wrong"})
        assert r.status_code == 200
    r = anon_client.post("/login", data={"email": "owner@example.com", "password": PASSWORD})
    assert r.status_code == 429


def test_open_redirect_blocked(app, anon_client):
    make_user()
    r = anon_client.post("/login?next=//evil.com", data={"email": "owner@example.com", "password": PASSWORD})
    assert r.headers["Location"] == "/"


def test_viewer_is_read_only(app):
    make_user("acct@example.com", role="viewer")
    c = app.test_client()
    c.post("/login", data={"email": "acct@example.com", "password": PASSWORD})
    assert c.get("/").status_code == 200
    assert c.post("/users/new", data={"email": "x@example.com", "password": "x" * 12}).status_code == 403


def test_password_change_signs_out_other_sessions(app, client):
    other = app.test_client()
    other.post("/login", data={"email": "owner@example.com", "password": PASSWORD})
    assert other.get("/").status_code == 200
    r = client.post("/account/password", data={"current": PASSWORD, "new": "another-long-pass", "confirm": "another-long-pass"})
    assert r.status_code == 302
    assert client.get("/").status_code == 200
    assert other.get("/").status_code == 302


def test_admin_creates_user(app, client):
    client.post("/users/new", data={"email": "partner@example.com", "role": "editor", "password": "x" * 12})
    assert User.query.filter_by(email="partner@example.com").one().role == "editor"


def test_csrf_enforced(app, client):
    app.config["WTF_CSRF_ENABLED"] = True
    assert client.post("/users/new", data={"email": "a@b.co", "password": "x" * 12}).status_code == 400


def test_full_export(app, client):
    import io

    from openpyxl import load_workbook

    resp = client.get("/export/all.xlsx")
    assert resp.status_code == 200
    wb = load_workbook(io.BytesIO(resp.data))
    assert "Mortgages" in wb.sheetnames and "Transactions" in wb.sheetnames
    assert "password_hash" not in [c.value for c in wb["Users"][1]]


def test_export_admin_only(app):
    make_user("viewer@example.com", role="viewer")
    c = app.test_client()
    c.post("/login", data={"email": "viewer@example.com", "password": PASSWORD})
    assert c.get("/export/all.xlsx").status_code == 403


def test_engine_options_handle_neon_pooler():
    from app import _engine_options

    pooled = _engine_options("postgresql://u:p@ep-cool-1-pooler.us-east-2.aws.neon.tech/neondb?sslmode=require")
    assert pooled["connect_args"] == {"prepare_threshold": None}
    direct = _engine_options("postgresql://u:p@ep-cool-1.us-east-2.aws.neon.tech/neondb?sslmode=require")
    assert "connect_args" not in direct and direct["pool_pre_ping"]


def test_admin_reset_from_environment(app, monkeypatch):
    from app import db
    from app.models import bootstrap_admin_from_env

    user = make_user("owner@example.com")
    user.active = False
    db.session.commit()
    monkeypatch.setenv("PMM_ADMIN_EMAIL", "Owner@example.com")
    monkeypatch.setenv("PMM_ADMIN_PASSWORD", "brand-new-password")
    bootstrap_admin_from_env()  # without the reset flag, existing users are left alone
    assert not db.session.get(User, user.id).check_password("brand-new-password")
    monkeypatch.setenv("PMM_ADMIN_RESET", "1")
    bootstrap_admin_from_env()
    u = db.session.get(User, user.id)
    assert u.check_password("brand-new-password") and u.active and u.role == "admin"
