import io
from datetime import date
from decimal import Decimal

from openpyxl import load_workbook

from app import db
from app.models import Mortgage
from app.services import importer

from .test_mortgages import create

HEADER = ",".join(h for h, _, _ in importer.COLUMNS)


def csv_with(*rows):
    lines = [HEADER]
    for r in rows:
        values = {h: "" for h, _, _ in importer.COLUMNS} | r
        lines.append(",".join(f'"{values[h]}"' for h, _, _ in importer.COLUMNS))
    return ("\n".join(lines) + "\n").encode()


def post(client, content, name="book.csv"):
    return client.post("/mortgages/import", data={"file": (io.BytesIO(content), name)}, content_type="multipart/form-data")


ROW = {"Borrower name": "Jane Smith", "Property address": "12 Maple Ave", "Principal": "250,000", "Rate %": "10.99",
       "Funded date": "01/02/2026", "Term months": "12", "Property type": "Semi-detached", "Rate type": "Fixed",
       "Compounding": "Semi-annual", "Payment type": "Interest only", "Payment frequency": "Bi-weekly",
       "Position": "2nd", "Bank keywords": "E-TRANSFER JANE; J SMITH", "Status": "Active"}


def test_template_downloads(client):
    resp = client.get("/mortgages/import/template.xlsx")
    wb = load_workbook(io.BytesIO(resp.data))
    assert wb["Mortgages"]["B1"].value == "Borrower name"
    assert client.get("/mortgages/import").status_code == 200


def test_import_csv_with_friendly_values(client):
    resp = post(client, csv_with(ROW, {**ROW, "Borrower name": "Bob Jones", "Reference": "B-7",
                                       "Current balance": "200000", "Balance as of": "2026-06-30"}))
    assert resp.status_code == 302
    jane = Mortgage.query.filter_by(borrower_name="Jane Smith").one()
    assert jane.reference == "M-001" and jane.property_type == "semi" and jane.compounding == "semi_annual"
    assert jane.payment_frequency == "biweekly" and jane.position == 2 and jane.funded_date == date(2026, 1, 2)
    assert jane.keywords == ["E-TRANSFER JANE", "J SMITH"]
    bob = Mortgage.query.filter_by(reference="B-7").one()
    assert bob.balance() == Decimal("200000.00") and bob.balance(as_of=date(2026, 6, 29)) == Decimal("250000.00")


def test_import_is_all_or_nothing(client):
    create(client)  # takes M-001
    bad = {**ROW, "Borrower name": "No Rate", "Rate %": ""}
    resp = post(client, csv_with(ROW, bad))
    assert resp.status_code == 400 and b"Row 3" in resp.data and b"Interest rate" in resp.data
    assert Mortgage.query.count() == 1
    # Duplicate reference also rejected.
    resp = post(client, csv_with({**ROW, "Reference": "M-001"}))
    assert resp.status_code == 400
    # Generated references skip existing ones.
    post(client, csv_with(ROW, {**ROW, "Borrower name": "Second"}))
    assert sorted(m.reference for m in Mortgage.query.all()) == ["M-001", "M-002", "M-003"]


def test_import_xlsx_template_ignores_example_row(client):
    wb = load_workbook(io.BytesIO(importer.template_workbook()))
    ws = wb["Mortgages"]
    headers = [c.value for c in ws[1]]
    row = [None] * len(headers)
    for h, v in {"Borrower name": "Excel Person", "Property address": "1 Bay St", "Principal": 300000, "Rate %": 9.5,
                 "Funded date": date(2026, 3, 1), "Maturity date": date(2027, 3, 1)}.items():
        row[headers.index(h)] = v
    ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    resp = post(client, buf.getvalue(), "book.xlsx")
    assert resp.status_code == 302
    assert [m.borrower_name for m in Mortgage.query.all()] == ["Excel Person"]
    assert db.session.query(Mortgage).one().term_months == 12
