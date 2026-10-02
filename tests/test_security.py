import io
from datetime import date
from decimal import Decimal

from openpyxl import load_workbook

from app import db
from app.models import BankTransaction, MortgageTransaction, StatementImport
from app.services import reports
from app.services.safety import csv_safe

from .test_mortgages import create

EVIL = '=HYPERLINK("http://evil.example","click")'


def test_csv_safe():
    assert csv_safe(EVIL).startswith("'=")
    assert csv_safe("-5 cash") == "'-5 cash"
    assert csv_safe("E-TRANSFER") == "E-TRANSFER"
    assert csv_safe(12) == 12


def test_exports_never_contain_live_formulas_from_bank_text(client):
    m = create(client, record_lender_fee="")
    imp = StatementImport(filename="x.csv", account_name="Chq")
    bt = BankTransaction(statement=imp, account_name="Chq", date=date(2026, 2, 1), description=EVIL,
                         amount=Decimal("1666.67"), fingerprint="f1", status="matched")
    db.session.add_all([imp, bt, MortgageTransaction(mortgage_id=m.id, date=date(2026, 2, 1), type="payment",
                                                     amount=Decimal("1666.67"), interest=Decimal("1666.67"),
                                                     principal=0, fees=0, bank_transaction=bt, notes=EVIL)])
    db.session.commit()
    r = reports.month_end_report(2026, 2)
    assert "'=HYPERLINK" in reports.transactions_csv(r)
    wb = load_workbook(io.BytesIO(reports.month_end_workbook(r)))
    cells = [c for row in wb["Transactions"].iter_rows() for c in row if c.value == EVIL]
    assert cells and all(c.data_type == "s" for c in cells)
    wb = load_workbook(io.BytesIO(client.get("/export/all.xlsx").data))
    cells = [c for row in wb["Bank lines"].iter_rows() for c in row if c.value == EVIL]
    assert cells and all(c.data_type == "s" for c in cells)


def test_security_headers_and_error_pages(client, anon_client):
    resp = client.get("/")
    assert "frame-ancestors 'none'" in resp.headers["Content-Security-Policy"]
    assert resp.headers["X-Frame-Options"] == "DENY"
    assert resp.headers["Cache-Control"] == "private, no-store"
    resp = client.get("/no-such-page")
    assert resp.status_code == 404 and b"That page doesn" in resp.data


def test_feed_parser_rejects_entity_expansion():
    import pytest
    from defusedxml import EntitiesForbidden

    from app.services import market

    bomb = b"""<?xml version="1.0"?><!DOCTYPE r [<!ENTITY a "aaaaaaaaaa"><!ENTITY b "&a;&a;&a;&a;&a;">]>
<rss><channel><item><title>&b;</title><link>https://x.example</link></item></channel></rss>"""
    with pytest.raises(EntitiesForbidden):
        market.parse_feed(bomb, "x")
