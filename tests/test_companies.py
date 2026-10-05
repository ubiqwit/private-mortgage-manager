"""Companies must never see each other's data."""
import io

from app import db
from app.models import BankTransaction, Company, Mortgage, MortgageDocument, MortgageTransaction
from app.tenancy import all_companies

from .conftest import PASSWORD, make_user
from .test_matching import confirm_mapping, upload
from .test_mortgages import BASE, create


def switch(client, name):
    company = Company.query.filter_by(name=name).one()
    assert client.post("/companies/switch", data={"company_id": company.id}).status_code == 302
    return company


def two_companies(client):
    """'My company' holds Jane Smith's mortgage; 'Beta Lending Inc.' holds Bob Jones's."""
    a = create(client)  # in the default company
    client.post("/companies/new", data={"name": "Beta Lending Inc.", "switch": "1"})
    b = create(client, borrower_name="Bob Jones", property_address="9 Oak St")
    return a.id, b.id


def test_lists_and_records_are_isolated(client):
    a_id, b_id = two_companies(client)
    page = client.get("/mortgages/?status=all").data
    assert b"Bob Jones" in page and b"Jane Smith" not in page
    assert client.get(f"/mortgages/{a_id}").status_code == 404  # other company's record by id
    assert client.get(f"/mortgages/{a_id}/schedule").status_code == 404
    assert client.post(f"/mortgages/{a_id}/transactions", data={"type": "payment", "date": "2026-02-01",
                                                                "amount": "100"}).status_code == 404
    assert client.post(f"/mortgages/{a_id}/delete", data={"confirm": "M-001"}).status_code == 404
    switch(client, "My company")
    page = client.get("/mortgages/?status=all").data
    assert b"Jane Smith" in page and b"Bob Jones" not in page
    dash = client.get("/").data
    assert b"Jane Smith" in dash or b"M-001" in dash
    assert db.session.execute(all_companies(db.select(Mortgage))).scalars().all().__len__() == 2


def test_references_restart_per_company(client):
    two_companies(client)
    refs = sorted(db.session.execute(all_companies(db.select(Mortgage.reference))).scalars())
    assert refs == ["M-001", "M-001"]  # unique within a company, not across companies


def test_reports_dashboard_and_statements_are_per_company(client):
    a_id, b_id = two_companies(client)  # now in Beta
    client.post(f"/mortgages/{b_id}/transactions", data={"type": "payment", "date": "2026-02-01", "amount": "1666.67"})
    feb = client.get("/reports/month-end?month=2026-02").data
    assert b"Bob Jones" in feb and b"Jane Smith" not in feb and b"Beta Lending Inc." in feb
    # The same statement file can be imported into both companies, and matches only its own mortgages.
    confirm_mapping(client, upload(client))
    switch(client, "My company")
    confirm_mapping(client, upload(client))
    lines = db.session.execute(all_companies(db.select(BankTransaction))).scalars().all()
    assert len(lines) == 8  # 4 lines × 2 companies, no false duplicates
    reconcile = client.get("/statements/reconcile?status=all&direction=all").data
    assert reconcile.count(b'title="E-TRANSFER AUTODEPOSIT BOB JONES"') == 1  # one row, this company's
    beta_line = next(b for b in lines if b.company_id != Company.query.filter_by(name="My company").one().id)
    assert client.post(f"/statements/lines/{beta_line.id}/ignore").status_code == 404
    # Matching a line to another company's mortgage is refused.
    mine = next(b for b in lines if b.company_id == Company.query.filter_by(name="My company").one().id)
    assert client.post(f"/statements/lines/{mine.id}/allocate", data={"mortgage_id": b_id}).status_code == 404


def test_documents_and_transactions_by_id_are_isolated(client):
    a_id, b_id = two_companies(client)  # in Beta
    switch(client, "My company")
    client.post(f"/mortgages/{a_id}/documents", data={"files": [(io.BytesIO(b"%PDF"), "a.pdf")], "category": "other"},
                content_type="multipart/form-data")
    client.post(f"/mortgages/{a_id}/transactions", data={"type": "payment", "date": "2026-02-01", "amount": "1666.67"})
    doc = db.session.execute(all_companies(db.select(MortgageDocument))).scalars().one()
    txn = db.session.execute(all_companies(db.select(MortgageTransaction).where(
        MortgageTransaction.type == "payment"))).scalars().one()
    switch(client, "Beta Lending Inc.")
    assert client.get(f"/mortgages/documents/{doc.id}/download").status_code == 404
    assert client.get(f"/mortgages/transactions/{txn.id}/edit").status_code == 404
    assert client.post(f"/mortgages/transactions/{txn.id}/delete").status_code == 404


def test_closed_months_are_per_company(client):
    two_companies(client)  # in Beta
    client.post("/reports/close?month=2026-02")
    assert b"Closed" in client.get("/reports/month-end?month=2026-02").data
    switch(client, "My company")
    assert b"Closed" not in client.get("/reports/month-end?month=2026-02").data


def test_users_only_reach_assigned_companies(app, client):
    two_companies(client)
    beta = Company.query.filter_by(name="Beta Lending Inc.").one()
    viewer = make_user("acct@example.com", role="viewer", companies=[beta])
    c = app.test_client()
    c.post("/login", data={"email": "acct@example.com", "password": PASSWORD})
    page = c.get("/mortgages/?status=all").data
    assert b"Bob Jones" in page and b"Jane Smith" not in page
    mine = Company.query.filter_by(name="My company").one()
    assert c.post("/companies/switch", data={"company_id": mine.id}).status_code == 403
    assert c.get("/companies/").status_code == 403
    # Admin removes all access → the viewer sees the "no company" page.
    client.post("/companies/access", data={})
    assert db.session.get(type(viewer), viewer.id).companies == []
    assert c.get("/").status_code == 403


def test_company_admin_create_rename_delete(client):
    two_companies(client)
    beta = Company.query.filter_by(name="Beta Lending Inc.").one()
    client.post(f"/companies/{beta.id}/rename", data={"name": "Beta Capital"})
    assert db.session.get(Company, beta.id).name == "Beta Capital"
    client.post(f"/companies/{beta.id}/delete")  # has a mortgage → refused
    assert db.session.get(Company, beta.id) is not None
    client.post("/companies/new", data={"name": "Empty Co"})
    empty = Company.query.filter_by(name="Empty Co").one()
    client.post(f"/companies/{empty.id}/delete")
    assert db.session.get(Company, empty.id) is None
    assert client.get("/companies/").status_code == 200


def test_demo_data_is_per_company(client):
    client.post("/companies/new", data={"name": "Sandbox", "switch": "1"})
    client.post("/demo/load")
    assert b"DEMO-01" in client.get("/mortgages/?status=all").data
    switch(client, "My company")
    assert b"DEMO-01" not in client.get("/mortgages/?status=all").data
    client.post("/demo/load")  # loads a separate demo set here too
    assert len(db.session.execute(all_companies(db.select(Mortgage))).scalars().all()) == 40


def test_full_export_covers_every_company(client):
    from openpyxl import load_workbook

    two_companies(client)
    wb = load_workbook(io.BytesIO(client.get("/export/all.xlsx").data))
    assert wb["Mortgages"].max_row == 3 and wb["Companies"].max_row == 3  # header + 2 rows each


def test_mortgage_created_in_current_company(client):
    client.post("/companies/new", data={"name": "Gamma", "switch": "1"})
    client.post("/mortgages/new", data=BASE)
    m = db.session.execute(all_companies(db.select(Mortgage))).scalars().one()
    assert m.company_id == Company.query.filter_by(name="Gamma").one().id
